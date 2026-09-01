from tkinter import filedialog, ttk

from customtkinter import CTkButton, CTkComboBox, CTkEntry, CTkFrame, CTkLabel, CTkToplevel
from openpyxl import load_workbook

from app.scanner import QRScanner

from .base_page import BasePage


class StudentsPage(BasePage):
    def __init__(self, master, controller):
        super().__init__(master, title="Students", sidebar=True)

        self.window = master
        self.controller = controller

        self.selected_student = None

        container = CTkFrame(master=self.wrapper, fg_color="transparent")
        container.pack(fill="x", padx=10, pady=10)

        # Make the 3 entry columns stretch horizontally
        container.grid_columnconfigure(0, weight=1)
        container.grid_columnconfigure(1, weight=1)
        container.grid_columnconfigure(2, weight=1)

        CTkLabel(master=container, text="Manage Students", text_color=("#464646", "#dcdcdc"), font=("Arial", 16, "bold")).grid(row=0, column=0, columnspan=3, padx=10, pady=10, sticky="w")

        self.student_first_name = CTkEntry(master=container, placeholder_text="First Name")
        self.student_first_name.grid(row=1, column=0, padx=(10, 5), pady=(10, 5), sticky="ew")

        self.student_last_name = CTkEntry(master=container, placeholder_text="Last Name")
        self.student_last_name.grid(row=1, column=1, padx=5, pady=(10, 5), sticky="ew")

        self.student_number = CTkEntry(master=container, placeholder_text="Student Number")
        self.student_number.grid(row=1, column=2, padx=(5, 10), pady=(10, 5), sticky="ew")

        self.course_values = []
        self.year_level_values = ["1st Year", "2nd Year", "3rd Year", "4th Year"]
        self.section_values = []

        self.course_options = CTkComboBox(master=container, values=[""], command=self.course_selected, button_color="#3B8ED0", button_hover_color="#4752AA")
        self.course_options.grid(row=2, column=0, padx=(10, 5), pady=(0, 5), sticky="ew")

        self.year_level = CTkComboBox(master=container, values=[""], button_color="#3B8ED0", button_hover_color="#4752AA", command=self.year_level_selected)
        self.year_level.grid(row=2, column=1, padx=5, pady=(0, 5), sticky="ew")

        self.section_options = CTkComboBox(master=container, values=[""], button_color="#3B8ED0", button_hover_color="#4752AA")
        self.section_options.grid(row=2, column=2, padx=(5, 10), pady=(0, 5), sticky="ew")

        CTkButton(master=container, text="Add Student", width=80, height=40, font=("", 16, "bold"), command=self.submit_form).grid(row=3, column=1, padx=10, pady=20)

        filter_controls = CTkFrame(master=self.wrapper, fg_color="transparent")
        filter_controls.pack(padx=10, pady=(0, 10), fill="x")

        self.student_search = CTkEntry(master=filter_controls, placeholder_text="Search student number or name...")
        self.student_search.pack(fill="x", padx=5, side="left", expand=True)
        self.student_search.bind("<KeyRelease>", self.search_students)

        CTkLabel(master=filter_controls, text="Course:").pack(side="left", padx=(10, 5))

        self.course_filter = CTkComboBox(master=filter_controls, values=["All Courses"], width=140, command=self.course_filter_selected)
        self.course_filter.pack(side="left", padx=5)

        CTkLabel(master=filter_controls, text="Year:").pack(side="left", padx=(10, 5))

        self.year_filter = CTkComboBox(master=filter_controls, values=["All Years"], width=120, command=self.year_filter_selected)
        self.year_filter.pack(side="left", padx=5)

        CTkLabel(master=filter_controls, text="Section:").pack(side="left", padx=(20, 5))

        self.section_filter = CTkComboBox(master=filter_controls, values=["All Sections"], width=160, command=self.section_filter_selected)
        self.section_filter.pack(side="left", padx=5)

        table_controls = CTkFrame(self.wrapper, fg_color="transparent")
        table_controls.pack(padx=10, pady=10, fill="x")

        CTkButton(master=table_controls, text="Import Excel", command=self.import_excel).pack(side="left", padx=(5, 5), pady=(0, 5))

        CTkButton(master=table_controls, text="Configure QR", command=self.open_scanner).pack(side="left", padx=(5, 5), pady=(0, 5))

        CTkButton(master=table_controls, text="Add to Class", command=self.open_enrollment).pack(side="left", padx=(5, 5), pady=(0, 5))

        CTkButton(master=table_controls, text="Delete Student", command=self.delete_student).pack(side="left", padx=(5, 5), pady=(0, 5))

        table_frame = CTkFrame(master=self.wrapper)
        table_frame.pack(fill="both", expand=True, padx=10, pady=10)

        columns = ("student_number", "name", "course", "year_level", "section")

        self.student_table = ttk.Treeview(table_frame, columns=columns, show="headings", selectmode="browse")

        self.student_table.heading("student_number", text="Student Number", anchor="w")
        self.student_table.heading("name", text="Student Name", anchor="w")
        self.student_table.heading("course", text="Course", anchor="w")
        self.student_table.heading("year_level", text="Year Level", anchor="w")
        self.student_table.heading("section", text="Section", anchor="w")

        self.student_table.column("student_number", width=130)
        self.student_table.column("name", width=200)
        self.student_table.column("course", width=100)
        self.student_table.column("year_level", width=100)
        self.student_table.column("section", width=130)

        # Vertical scrollbar
        scrollbar = ttk.Scrollbar(table_frame, orient="vertical", command=self.student_table.yview)

        self.student_table.configure(yscrollcommand=scrollbar.set)

        self.student_table.pack(side="left", fill="both", expand=True)

        scrollbar.pack(side="right", fill="y")

        style = ttk.Style()
        style.configure("Treeview.Heading", font=("", 20, "bold"))
        style.configure("Treeview", font=("", 16))

        self.student_table.bind("<<TreeviewSelect>>", self.student_selected)

    def search_students(self, event=None):
        self.display_students()

    def delete_student(self):
        if not self.selected_student:
            self.window.show_error(title="No Student Selected", message="Select a student first.")
            return

        student = self.selected_student
        confirm = self.window.ask_yes_no(
            title="Delete student", message=f"Are you sure you want to remove:\n\n{student.student_number} - {student.first_name} {student.last_name}?", parent=self.window
        )

        if confirm is not True:
            return

        deleted = self.controller.delete_student(student.id)

        if not deleted:
            self.window.show_error(title="Delete Student", message="Student could not be deleted.")
            return

        self.selected_student = None
        self.refresh()

        self.window.show_info(title="Delete Student", message=(f"{student.student_number} - {student.first_name} {student.last_name} was deleted successfully."))

    def open_enrollment(self):
        if not self.selected_student:
            self.window.show_error(title="No Student Selected", message="Select a student first.")
            return

        student = self.selected_student

        classes = self.controller.get_classes_by_section(student.section_id)

        if not classes:
            self.window.show_error(title="No Classes", message="There are no classes available for this student's section.")
            return

        self.enrollment_classes = classes

        self.enrollment_window = CTkToplevel(self.master)
        self.enrollment_window.title("Add Student to Class")
        self.enrollment_window.geometry("400x250")
        self.enrollment_window.resizable(False, False)

        self.enrollment_window.grab_set()
        self.enrollment_window.focus_force()

        CTkLabel(self.enrollment_window, text="Enroll Student", font=("Arial", 18, "bold")).pack(pady=(20, 5))

        CTkLabel(self.enrollment_window, text=(f"{student.student_number} - {student.first_name} {student.last_name}")).pack(pady=(0, 15))

        class_values = []

        for class_ in classes:
            subject = class_.subject

            if subject:
                enrolled = self.controller.is_student_enrolled(student.id, class_.id)

                status = " (Enrolled)" if enrolled else ""

                class_values.append(f"{subject.name}{status}")

        if not class_values:
            self.enrollment_window.destroy()

            self.window.show_error(title="No Subjects", message="No subjects are available for this student's section.")
            return

        self.class_options = CTkComboBox(self.enrollment_window, values=class_values, width=300)
        self.class_options.pack(pady=10)

        self.class_options.set("Select subject")

        CTkButton(
            self.enrollment_window,
            text="Enroll",
            command=self.enroll_student,  # type: ignore
        ).pack(pady=15)

    def enroll_student(self):
        if not self.selected_student:
            return

        selected = self.class_options.get().strip()

        if selected == "Select subject" or not selected:
            self.window.show_error(title="Enrollment Error", message="Select a subject first.")
            return

        # Remove the "(Enrolled)" display suffix.
        subject_name = selected.replace(" (Enrolled)", "")

        selected_class = None

        for class_ in self.enrollment_classes:
            if class_.subject and class_.subject.name == subject_name:
                selected_class = class_
                break

        if not selected_class:
            self.window.show_error(title="Enrollment Error", message="The selected subject could not be found.")
            return

        if self.controller.is_student_enrolled(self.selected_student.id, selected_class.id):
            self.window.show_error(title="Enrollment", message="Student is already enrolled in this subject.")
            return

        enrollment = self.controller.enroll_student(student_id=self.selected_student.id, class_id=selected_class.id)

        if enrollment:
            self.window.show_info(title="Enrollment", message=(f"{self.selected_student.first_name} {self.selected_student.last_name} was enrolled in {subject_name}."))

            self.enrollment_window.destroy()

    def course_filter_selected(self, course_name):
        self.year_filter.set("All Years")

        self.section_filter.configure(values=["All Sections"])
        self.section_filter.set("All Sections")

        self.display_students()

    def year_filter_selected(self, year_level):
        course_name = self.course_filter.get().strip()

        if course_name == "All Courses":
            self.section_filter.configure(values=["All Sections"])
            self.section_filter.set("All Sections")
            self.display_students()
            return

        course = self.controller.get_course(course_name)

        if not course:
            return

        if year_level == "All Years":
            section_values = ["All Sections"]
        else:
            sections = self.controller.get_sections_by_course_and_year(course.id, year_level)

            section_values = ["All Sections"] + [section.name for section in sections]

        self.section_filter.configure(values=section_values)
        self.section_filter.set("All Sections")

        self.display_students()

    def section_filter_selected(self, section_name):
        self.display_students()

    def year_level_selected(self, year_level, event=None):
        course_name = self.course_options.get().strip().upper()
        year_level = year_level.strip()

        course = self.controller.get_course(course_name)

        if not course or not year_level:
            self.section_options.configure(values=[])
            self.section_options.set("Select section")
            return

        sections = self.controller.get_sections_by_course_and_year(course.id, year_level)

        self.section_values = [section.name for section in sections]

        self.section_options.configure(values=self.section_values)
        self.section_options.set("Select section")

    def course_selected(self, course, event=None):
        course_name = course.strip().upper()
        existing_course = self.controller.get_course(course_name)

        if not existing_course:
            self.section_options.configure(values=[])
            self.section_options.set("Select section")
            return

        year_level = self.year_level.get().strip()

        if not year_level or year_level == "Select year level":
            self.section_options.configure(values=[])
            self.section_options.set("Select section")
            return

        sections = self.controller.get_sections_by_course_and_year(existing_course.id, year_level)

        self.section_values = [section.name for section in sections]

        self.section_options.configure(values=self.section_values)
        self.section_options.set("Select section")

    def submit_form(self, event=None):
        st_num = self.student_number.get().strip().upper()
        if not st_num:
            self.window.show_error(title="Student error", message="Student number is required.")
            return

        course_name = self.course_options.get().strip().upper()
        section_name = self.section_options.get().strip().upper()
        course = self.controller.get_course(course_name)
        year_level = self.year_level.get()

        if not course:
            course = self.controller.add_course(course_name)

        section = self.controller.get_section_by_course_and_name(course.id, section_name)

        if not section:
            section = self.controller.add_section(name=section_name, course_id=course.id, year_level=year_level)

        student = self.controller.add_student(
            first_name=self.student_first_name.get().title(), last_name=self.student_last_name.get().title(), student_number=self.student_number.get().strip().upper(), section_id=section.id
        )

        if student is None:
            self.window.show_error(title="Student error", message="Student already exists.")

        else:
            self.window.show_info(title="Student Management", message="Added student to database.")
            self.clear_form()
            self.refresh()

    def clear_form(self):
        self.student_first_name.delete(0, "end")
        self.student_last_name.delete(0, "end")
        self.student_number.delete(0, "end")
        self.course_options.set("Select course")
        self.year_level.set("Select year level")
        self.section_options.set("Select section")

    def display_students(self):
        for item in self.student_table.get_children():
            self.student_table.delete(item)

        course_filter = self.course_filter.get().strip()
        year_filter = self.year_filter.get().strip()
        section_filter = self.section_filter.get().strip()
        search_text = self.student_search.get().strip().lower()

        students = self.controller.get_all_students_with_section()

        for student in students:
            section = student.section
            course = section.course

            full_name = f"{student.first_name} {student.last_name}".lower()
            student_number = student.student_number.lower()

            if search_text:
                if search_text not in full_name and search_text not in student_number:
                    continue

            if course_filter != "All Courses":
                if course.name != course_filter:
                    continue

            if year_filter != "All Years":
                if section.year_level != year_filter:
                    continue

            if section_filter != "All Sections":
                if section.name != section_filter:
                    continue

            self.student_table.insert(
                "", "end", iid=str(student.id), values=(student.student_number, f"{student.first_name} {student.last_name}", course.name, section.year_level, section.name)
            )

        self.selected_student = None

    def student_selected(self, event=None):
        selection = self.student_table.selection()

        if not selection:
            self.selected_student = None
            return

        student_id = int(selection[0])

        self.selected_student = self.controller.get_student(student_id)

    def import_excel(self):
        file_path = filedialog.askopenfilename(title="Select Excel File", filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")])

        if not file_path:
            return

        workbook = load_workbook(file_path, read_only=True, data_only=True)
        sheet = workbook.active

        assert sheet is not None
        headers = [cell.value for cell in next(sheet.iter_rows())]

        required_headers = ["student_number", "first_name", "last_name", "course", "year_level", "section"]

        if not all(header in headers for header in required_headers):
            workbook.close()
            self.window.show_error(title="Import Error", message=("Invalid Excel format.\n\nRequired columns:\nstudent_number, first_name, last_name, course, year_level, section"))
            return

        header_index = {header: index for index, header in enumerate(headers)}

        imported = 0
        skipped = 0

        assert sheet is not None
        for row in sheet.iter_rows(min_row=2, values_only=True):
            student_number = row[header_index["student_number"]]

            if not student_number:
                skipped += 1
                continue

            student_number = str(student_number).strip().upper()

            first_name = str(row[header_index["first_name"]] or "").strip()

            last_name = str(row[header_index["last_name"]] or "").strip()

            course_name = str(row[header_index["course"]] or "").strip().upper()

            year_level = str(row[header_index["year_level"]] or "").strip()

            section_name = str(row[header_index["section"]] or "").strip().upper()

            if not all([first_name, last_name, course_name, year_level, section_name]):
                skipped += 1
                continue

            course = self.controller.get_course(course_name)

            if not course:
                course = self.controller.add_course(course_name)

            section = self.controller.get_section_by_course_and_name(course.id, section_name)

            if not section:
                section = self.controller.add_section(name=section_name, course_id=course.id, year_level=year_level)

            student = self.controller.add_student(first_name=first_name.title(), last_name=last_name.title(), student_number=student_number, section_id=section.id)

            if student is None:
                skipped += 1
            else:
                imported += 1

        workbook.close()

        self.refresh()

        self.window.show_info(title="Import Students", message=(f"Import complete.\n\nImported: {imported}\nSkipped: {skipped}"))

    def open_scanner(self, event=None):
        if not self.selected_student:
            self.window.show_error(title="No Student Selected", message="Select a student first.")
            return

        self.scanner = QRScanner(master=self.master, controller=self.controller, on_scan=self.register_qr)

    def register_qr(self, qr_data):
        student = self.selected_student

        if not student:
            return

        qr = self.controller.configure_qr(student_id=student.id, qr_hash=qr_data)

        if not qr:
            self.window.show_error(title="QR Code Error", message="QR is already taken.")
        elif student.qr_hash == qr_data:
            self.window.show_info(title="QR Registration", message=f"QR already registered to {student.student_number} - {student.last_name}.")
        else:
            self.window.show_info(title="QR Registration", message=f"Successfully registered to {student.student_number}.")

    def refresh(self):
        courses = self.controller.get_all_courses()

        self.course_values = [course.name for course in courses]

        # Form
        self.course_options.configure(values=self.course_values)
        self.course_options.set("Select course")

        self.year_level.configure(values=self.year_level_values)
        self.year_level.set("Select year level")

        self.section_options.configure(values=[])
        self.section_options.set("Select section")

        # Table filters
        self.course_filter.configure(values=["All Courses"] + self.course_values)
        self.course_filter.set("All Courses")

        self.year_filter.configure(values=["All Years"] + self.year_level_values)
        self.year_filter.set("All Years")

        self.section_filter.configure(values=["All Sections"])
        self.section_filter.set("All Sections")

        self.display_students()
